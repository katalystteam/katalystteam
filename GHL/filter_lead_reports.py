"""Filter Crexi lead report workbooks: filter/sort Lead Report + Detail only,
leave all other sheets untouched, save as *__cleaned.xlsx."""

from __future__ import annotations

from copy import copy
from datetime import date, datetime
from pathlib import Path
import shutil

import openpyxl

CUTOFF = datetime(2026, 6, 15)
PROCESS_SHEETS = ("1 - Lead Report", "1 - Detail")
DATE_HEADERS = {
    "1 - Lead Report": "Last Action Date",
    "1 - Detail": "Date",
}

FILES = [
    Path("/Users/john.angeles/Downloads/Lead_Report_Chatham_Ave_Land.xlsx"),
    Path("/Users/john.angeles/Downloads/Lead_Report_2733_2897_Boston_Ave_Des_Moines_IA.xlsx"),
    Path("/Users/john.angeles/Downloads/Lead_Report_811_16th_St_Des_Moines_IA.xlsx"),
    Path("/Users/john.angeles/Downloads/Lead_Report_405_N_Davis_St.xlsx"),
    Path("/Users/john.angeles/Downloads/Lead_Report_26th_Street_Apartments.xlsx"),
]


def _as_datetime(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    parsed = openpyxl.utils.datetime.from_excel(value) if isinstance(value, (int, float)) else None
    if isinstance(parsed, datetime):
        return parsed
    if isinstance(parsed, date):
        return datetime.combine(parsed, datetime.min.time())
    return None


def _find_date_column(ws, header_name: str) -> tuple[int | None, int | None]:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        for cell in row:
            if cell.value == header_name:
                return cell.row, cell.column
    return None, None


def _row_has_data(ws, row_idx: int, name_col: int = 2) -> bool:
    value = ws.cell(row=row_idx, column=name_col).value
    return value is not None and str(value).strip() != ""


def _process_sheet(ws, sheet_name: str) -> tuple[int, int]:
    header_name = DATE_HEADERS[sheet_name]
    header_row, date_col = _find_date_column(ws, header_name)
    if header_row is None or date_col is None:
        raise ValueError(f"{sheet_name}: could not find {header_name!r} column")

    data_start = header_row + 1
    data_end = data_start
    while data_end <= ws.max_row and _row_has_data(ws, data_end):
        data_end += 1
    data_end -= 1

    kept_rows: list[tuple[datetime, int]] = []
    for row_idx in range(data_start, data_end + 1):
        dt = _as_datetime(ws.cell(row=row_idx, column=date_col).value)
        if dt is not None and dt >= CUTOFF:
            kept_rows.append((dt, row_idx))

    kept_rows.sort(key=lambda item: item[0], reverse=True)
    original_count = max(0, data_end - data_start + 1)
    max_col = ws.max_column

    saved_rows: list[list[tuple[object, object]]] = []
    for _, src_row in kept_rows:
        saved_rows.append(
            [
                (
                    ws.cell(row=src_row, column=col).value,
                    copy(ws.cell(row=src_row, column=col)._style),
                )
                for col in range(1, max_col + 1)
            ]
        )

    if data_end >= data_start:
        ws.delete_rows(data_start, data_end - data_start + 1)

    for offset, row_cells in enumerate(saved_rows):
        target_row = data_start + offset
        for col, (value, style) in enumerate(row_cells, start=1):
            cell = ws.cell(row=target_row, column=col, value=value)
            cell._style = style

    return original_count, len(saved_rows)


def cleaned_output_path(path: Path) -> Path:
    return path.with_name(f"{path.stem}__cleaned{path.suffix}")


def process_workbook(path: Path) -> Path:
    output_path = cleaned_output_path(path)
    shutil.copy2(path, output_path)
    wb = openpyxl.load_workbook(output_path)

    summary: list[str] = []
    for sheet_name in PROCESS_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        before, after = _process_sheet(wb[sheet_name], sheet_name)
        summary.append(f"{sheet_name}: {before} -> {after} rows")

    wb.save(output_path)
    print(f"{path.name} -> {output_path.name}")
    print(f"  sheets kept: {', '.join(wb.sheetnames)}")
    for line in summary:
        print(f"  {line}")
    return output_path


def main() -> None:
    for path in FILES:
        if not path.exists():
            raise FileNotFoundError(path)
        process_workbook(path)


if __name__ == "__main__":
    main()
