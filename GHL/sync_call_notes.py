"""
Sync call notes from Excel (Economics sheets) to GoHighLevel contact notes.

Reads the "Want to Meet With / Research" section from Economics sheets in
GHL/jared_call_notes.xlsx (manual export from Google Sheets), finds contacts in GHL by name, and creates notes on the
matching contact. Supports a full reset: purge existing [Call Notes - ...] notes
then re-create oldest → newest so the newest call date appears at the top in GHL.

API Reference (official):
  - Search Contacts: GET /contacts/?locationId=...&query=...
  - Get Notes: GET /contacts/{contactId}/notes
  - Create Note: POST /contacts/{contactId}/notes  (body: { "body": "...", "userId": "..." })
  - Delete Note: DELETE /contacts/{contactId}/notes/{noteId}
  - Base URL: https://services.leadconnectorhq.com
  - Header: Version: 2021-07-28
  - Auth: Bearer <Private Integration Token or OAuth Access Token>
  - Scopes needed: contacts.readonly, contacts.write
"""

from __future__ import annotations

import csv
import os
import re
import sys
import shutil
import logging
import time
from dataclasses import dataclass, field
from datetime import datetime

import requests
import openpyxl

try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env.local"))
    load_dotenv(os.path.join(os.path.dirname(__file__), "..", "real-estate-crm-dashboard-react", ".env.local"))
except ImportError:
    pass

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
GHL_BASE_URL = "https://services.leadconnectorhq.com"
GHL_API_VERSION = "2021-07-28"

# Set these via environment variables (never hardcode secrets)
GHL_API_TOKEN = os.environ.get("GHL_API_TOKEN", "")
GHL_LOCATION_ID = os.environ.get("GHL_LOCATION_ID", "")

DEFAULT_WORKBOOK = os.path.join(os.path.dirname(__file__), "jared_call_notes.xlsx")
EXCEL_PATH = os.environ.get("CALL_NOTES_FILE", DEFAULT_WORKBOOK)
LOG_DIR = os.path.join(os.path.dirname(__file__), "..", "sync_logs")
SYNC_NOTE_PREFIX = "[Call Notes - "
ECONOMICS_TEMPLATE_SHEET = "Economics Template"

# Rate-limit safety: pause between API calls (seconds)
REQUEST_DELAY = 0.5

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------
@dataclass
class CallNote:
    name: str
    notes: str
    called: bool
    sheet_date: str  # e.g. "1.5.26"
    phone: str = ""  # phone number extracted from the name column, if any


@dataclass
class SyncLogEntry:
    name: str
    sheet_date: str
    status: str  # synced, not_found, skipped_duplicate, failed
    note_body: str = ""
    ghl_contact_id: str = ""
    timestamp: str = field(default_factory=lambda: datetime.now().isoformat())


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------
def resolve_workbook_path(path: str | None = None) -> str:
    """Return a path openpyxl can read (.xlsx). Copies extensionless exports if needed."""
    if path and os.path.isfile(path):
        source = os.path.abspath(path)
    elif os.environ.get("CALL_NOTES_FILE", "").strip():
        source = os.path.abspath(os.environ["CALL_NOTES_FILE"].strip())
    else:
        source = os.path.abspath(EXCEL_PATH)

    if not os.path.isfile(source):
        raise FileNotFoundError(
            f"Call notes workbook not found: {source}\n"
            "Download the Google Sheet export to GHL/jared_call_notes.xlsx"
        )

    if source.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        return source

    resolved = f"{source}.xlsx"
    if not os.path.isfile(resolved) or os.path.getmtime(resolved) < os.path.getmtime(source):
        shutil.copy2(source, resolved)
    return resolved


def parse_name_and_phone(raw_name: str) -> tuple:
    """Split a raw name like 'Keith Kleinhans 281-777-7177' into (name, phone)."""
    # Match common US phone patterns: 281-777-7177, (281) 777-7177, 281.777.7177, etc.
    phone_pattern = r'[\(]?\d{3}[\)\-\.\s]?\s?\d{3}[\-\.\s]?\d{4}'
    match = re.search(phone_pattern, raw_name)
    if match:
        phone = match.group(0).strip()
        name = raw_name[:match.start()].strip().rstrip('-').strip()
        return name, phone
    return raw_name.strip(), ""


def parse_sheet_date(sheet_name: str) -> str:
    """Extract the date portion from a sheet name like '1.5.26 Economics'."""
    return sheet_name.replace(" Economics", "").strip()


def get_economics_sheets(excel_path: str) -> list[str]:
    """Return Economics sheet names from the workbook (newest tab first, per workbook order)."""
    workbook = resolve_workbook_path(excel_path)
    wb = openpyxl.load_workbook(workbook, read_only=True)
    sheets = [
        s for s in wb.sheetnames
        if "Economics" in s and s != ECONOMICS_TEMPLATE_SHEET
    ]
    wb.close()
    return sheets


def get_economics_sheets_chronological(excel_path: str) -> list[str]:
    """Return Economics sheets oldest → newest (workbook lists newest tabs first)."""
    return list(reversed(get_economics_sheets(excel_path)))


def extract_all_call_notes_chronological(excel_path: str) -> list[CallNote]:
    """Extract call notes from all Economics sheets in oldest → newest order."""
    workbook = resolve_workbook_path(excel_path)
    sheet_names = get_economics_sheets_chronological(workbook)
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ordered: list[CallNote] = []
    try:
        for sheet_name in sheet_names:
            ordered.extend(_extract_call_notes_from_ws(wb[sheet_name], sheet_name))
    finally:
        wb.close()
    logger.info(f"Extracted {len(ordered)} total call notes (chronological)")
    return ordered


def _extract_call_notes_from_ws(ws, sheet_name: str) -> list[CallNote]:
    """Extract call notes from an already-open worksheet."""
    sheet_date = parse_sheet_date(sheet_name)
    header_found = False
    results: list[CallNote] = []

    for row in ws.iter_rows(values_only=True):
        if not header_found:
            if any("Want to Meet" in str(v) for v in row if v is not None):
                header_found = True
            continue

        if len(row) < 4:
            continue

        name = row[1]
        called = row[2]
        notes = row[3]

        if name is None or str(name).strip() == "":
            continue

        raw_name = str(name).strip()
        clean_name, phone = parse_name_and_phone(raw_name)

        if called and notes and str(notes).strip():
            results.append(CallNote(
                name=clean_name,
                notes=str(notes).strip(),
                called=bool(called),
                sheet_date=sheet_date,
                phone=phone,
            ))

    logger.info(f"Extracted {len(results)} call notes from sheet '{sheet_name}'")
    return results


def extract_call_notes(excel_path: str, sheet_name: str) -> list[CallNote]:
    """
    Read the specified Economics sheet and return CallNote entries
    from the 'Want to Meet With / Research' section where Called=True
    and Notes/Tasks is not empty.
    """
    workbook = resolve_workbook_path(excel_path)
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames[:10]}...")

    try:
        return _extract_call_notes_from_ws(wb[sheet_name], sheet_name)
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# GHL API helpers
# ---------------------------------------------------------------------------
def _headers() -> dict:
    return {
        "Authorization": f"Bearer {GHL_API_TOKEN}",
        "Content-Type": "application/json",
        "Version": GHL_API_VERSION,
    }


def search_contact_by_name(name: str) -> dict | None:
    """
    Search for a GHL contact by name using GET /contacts/.
    Returns the first matching contact dict or None.

    Per the official API spec (contacts.json from GoHighLevel/highlevel-api-docs):
      GET /contacts/
      Query params: locationId (required), query (optional, searches by name)
    """
    url = f"{GHL_BASE_URL}/contacts/"
    params = {
        "locationId": GHL_LOCATION_ID,
        "query": name,
        "limit": 5,
    }

    try:
        resp = requests.get(url, headers=_headers(), params=params, timeout=30)
        resp.raise_for_status()
        data = resp.json()
    except requests.RequestException as e:
        logger.error(f"Failed to search for contact '{name}': {e}")
        return None

    contacts = data.get("contacts", [])
    if not contacts:
        logger.warning(f"No GHL contact found for '{name}'")
        return None

    # Try exact match first (case-insensitive)
    name_lower = name.lower()
    for c in contacts:
        full_name = f"{c.get('firstName', '')} {c.get('lastName', '')}".strip().lower()
        contact_name = c.get("name", "").lower() if c.get("name") else ""
        if full_name == name_lower or contact_name == name_lower:
            logger.info(f"Exact match found for '{name}': contactId={c['id']}")
            return c

    # Fall back to first result if no exact match
    first = contacts[0]
    first_name = f"{first.get('firstName', '')} {first.get('lastName', '')}".strip()
    logger.info(f"Using closest match for '{name}': '{first_name}' (contactId={first['id']})")
    return first


def create_note(contact_id: str, note_body: str) -> dict | None:
    """
    Create a note on a GHL contact.

    Per the official API spec (contacts.json from GoHighLevel/highlevel-api-docs):
      POST /contacts/{contactId}/notes
      Body (NotesDTO): { "body": "<string>" (required), "userId": "<string>" (optional) }
      Returns 201 with { "note": { "id", "body", "userId", "dateAdded", "contactId" } }
    """
    url = f"{GHL_BASE_URL}/contacts/{contact_id}/notes"
    payload = {"body": note_body}

    try:
        resp = requests.post(url, headers=_headers(), json=payload, timeout=30)
        resp.raise_for_status()
        data = resp.json()
        note = data.get("note", {})
        logger.info(f"Note created: noteId={note.get('id')} for contactId={contact_id}")
        return note
    except requests.RequestException as e:
        logger.error(f"Failed to create note for contactId={contact_id}: {e}")
        if hasattr(e, "response") and e.response is not None:
            logger.error(f"Response body: {e.response.text}")
        return None


def get_existing_notes(contact_id: str) -> list[dict]:
    """
    Get all existing notes for a contact.

    Per the official API spec:
      GET /contacts/{contactId}/notes
      Returns { "notes": [ { "id", "body", "userId", "dateAdded", "contactId" } ] }
    """
    url = f"{GHL_BASE_URL}/contacts/{contact_id}/notes"

    try:
        resp = requests.get(url, headers=_headers(), timeout=30)
        resp.raise_for_status()
        data = resp.json()
        return data.get("notes", [])
    except requests.RequestException as e:
        logger.error(f"Failed to get notes for contactId={contact_id}: {e}")
        return []


def note_already_exists(contact_id: str, note_body: str) -> bool:
    """Check if a note with the same body text already exists on the contact."""
    existing = get_existing_notes(contact_id)
    for note in existing:
        if note.get("body", "").strip() == note_body.strip():
            return True
    return False


def is_synced_call_note(body: str) -> bool:
    """True for notes created by this sync script."""
    return body.strip().startswith(SYNC_NOTE_PREFIX)


def delete_note(contact_id: str, note_id: str) -> bool:
    """
    Delete a note from a GHL contact.

    DELETE /contacts/{contactId}/notes/{id}
    """
    url = f"{GHL_BASE_URL}/contacts/{contact_id}/notes/{note_id}"

    try:
        resp = requests.delete(url, headers=_headers(), timeout=30)
        resp.raise_for_status()
        logger.info(f"Note deleted: noteId={note_id} for contactId={contact_id}")
        return True
    except requests.RequestException as e:
        logger.error(f"Failed to delete noteId={note_id} for contactId={contact_id}: {e}")
        if hasattr(e, "response") and e.response is not None:
            logger.error(f"Response body: {e.response.text}")
        return False


def purge_synced_notes(contact_id: str, dry_run: bool = False) -> dict:
    """Delete all [Call Notes - ...] notes on a contact. Returns {deleted, failed, matched}."""
    summary = {"deleted": 0, "failed": 0, "matched": 0}
    for note in get_existing_notes(contact_id):
        body = note.get("body", "")
        if not is_synced_call_note(body):
            continue
        summary["matched"] += 1
        if dry_run:
            summary["deleted"] += 1
            continue
        if delete_note(contact_id, note["id"]):
            summary["deleted"] += 1
        else:
            summary["failed"] += 1
        time.sleep(REQUEST_DELAY)
    return summary


def _require_ghl_env() -> None:
    if not GHL_API_TOKEN:
        logger.error("GHL_API_TOKEN environment variable is not set")
        sys.exit(1)
    if not GHL_LOCATION_ID:
        logger.error("GHL_LOCATION_ID environment variable is not set")
        sys.exit(1)


def _lookup_contacts(names: list[str]) -> dict[str, dict | None]:
    """Search GHL once per unique contact name."""
    cache: dict[str, dict | None] = {}
    for i, name in enumerate(names, 1):
        if name in cache:
            continue
        logger.info(f"[lookup {i}/{len(names)}] {name}")
        cache[name] = search_contact_by_name(name)
        time.sleep(REQUEST_DELAY)
    return cache


def reset_call_notes(
    excel_path: str | None = None,
    dry_run: bool = False,
    write_log: bool = True,
) -> dict:
    """
    Purge existing synced call notes, then recreate all notes oldest → newest.

    Creates notes in chronological order so the newest call date appears at the
    top of each contact's notes list in GHL.
    """
    _require_ghl_env()
    source = excel_path or EXCEL_PATH
    workbook = resolve_workbook_path(source)
    logger.info(f"Reading from: {workbook}")
    logger.info(f"Dry run: {dry_run}")

    call_notes = extract_all_call_notes_chronological(workbook)
    unique_names = list(dict.fromkeys(cn.name for cn in call_notes))
    logger.info(f"Unique contacts: {len(unique_names)}")

    contact_by_name = _lookup_contacts(unique_names)

    totals = {
        "total": len(call_notes),
        "purge_matched": 0,
        "purge_deleted": 0,
        "purge_failed": 0,
        "synced": 0,
        "failed": 0,
        "not_found": 0,
        "not_found_contacts": [],
        "log_entries": [],
        "sheets_processed": len(get_economics_sheets_chronological(workbook)),
    }

    # Phase 1: purge synced notes on matched contacts
    logger.info("=" * 60)
    logger.info("Phase 1: Purge existing [Call Notes - ...] notes")
    logger.info("=" * 60)
    purged_contacts: set[str] = set()
    for name, contact in contact_by_name.items():
        if not contact:
            continue
        contact_id = contact["id"]
        if contact_id in purged_contacts:
            continue
        purged_contacts.add(contact_id)
        purge = purge_synced_notes(contact_id, dry_run=dry_run)
        totals["purge_matched"] += purge["matched"]
        totals["purge_deleted"] += purge["deleted"]
        totals["purge_failed"] += purge["failed"]

    # Phase 2: create notes oldest → newest
    logger.info("=" * 60)
    logger.info("Phase 2: Create call notes (oldest → newest)")
    logger.info("=" * 60)
    for i, cn in enumerate(call_notes, 1):
        note_body = format_note_body(cn)
        logger.info(f"[{i}/{len(call_notes)}] {cn.name} ({cn.sheet_date})")

        contact = contact_by_name.get(cn.name)
        if not contact:
            totals["not_found"] += 1
            totals["not_found_contacts"].append((cn.name, cn.sheet_date))
            totals["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="not_found", note_body=note_body,
            ))
            continue

        contact_id = contact["id"]
        if dry_run:
            totals["synced"] += 1
            totals["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="synced (dry_run)", note_body=note_body,
                ghl_contact_id=contact_id,
            ))
            continue

        result = create_note(contact_id, note_body)
        if result:
            totals["synced"] += 1
            totals["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="synced", note_body=note_body,
                ghl_contact_id=contact_id,
            ))
        else:
            totals["failed"] += 1
            totals["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="failed", note_body=note_body,
                ghl_contact_id=contact_id,
            ))
        time.sleep(REQUEST_DELAY)

    _print_reset_summary(totals)
    if write_log and totals["log_entries"]:
        label = "reset_dry_run" if dry_run else "reset"
        write_sync_log(totals["log_entries"], label=label)
    return totals


def _print_reset_summary(summary: dict) -> None:
    logger.info("*" * 60)
    logger.info("Reset Summary:")
    logger.info(f"  Economics sheets:    {summary['sheets_processed']}")
    logger.info(f"  Call notes total:    {summary['total']}")
    logger.info(f"  Purge matched:       {summary['purge_matched']}")
    logger.info(f"  Purge deleted:       {summary['purge_deleted']}")
    logger.info(f"  Purge failed:        {summary['purge_failed']}")
    logger.info(f"  Created:             {summary['synced']}")
    logger.info(f"  Not found in GHL:    {summary['not_found']}")
    logger.info(f"  Create failed:       {summary['failed']}")
    if summary["not_found_contacts"]:
        not_found_map: dict[str, list[str]] = {}
        for name, sheet_date in summary["not_found_contacts"]:
            not_found_map.setdefault(name, []).append(sheet_date)
        logger.info("")
        logger.info(f"  Contacts NOT found in GHL ({len(not_found_map)} unique):")
        for name in sorted(not_found_map.keys()):
            dates = ", ".join(not_found_map[name])
            logger.info(f"    - {name}  (sheets: {dates})")
    logger.info("*" * 60)


# ---------------------------------------------------------------------------
# Main sync logic
# ---------------------------------------------------------------------------
def format_note_body(call_note: CallNote) -> str:
    """Format the note body with date context and phone if available."""
    parts = [f"[Call Notes - {call_note.sheet_date}]"]
    if call_note.phone:
        parts.append(f"(Phone: {call_note.phone})")
    parts.append(call_note.notes)
    return " ".join(parts)


def sync_notes(
    sheet_name: str,
    dry_run: bool = False,
    write_log: bool = True,
    excel_path: str | None = None,
) -> dict:
    """
    Main function: extract call notes from Excel and sync to GHL.

    Args:
        sheet_name: Name of the Economics sheet (e.g., "1.5.26 Economics")
        dry_run: If True, only log what would happen without making API calls

    Returns:
        Summary dict with counts of synced, skipped, failed, not_found.
    """
    if not GHL_API_TOKEN:
        logger.error("GHL_API_TOKEN environment variable is not set")
        sys.exit(1)
    if not GHL_LOCATION_ID:
        logger.error("GHL_LOCATION_ID environment variable is not set")
        sys.exit(1)

    workbook = resolve_workbook_path(excel_path or EXCEL_PATH)
    logger.info(f"Reading from: {workbook}")
    logger.info(f"Sheet: {sheet_name}")
    logger.info(f"Dry run: {dry_run}")

    call_notes = extract_call_notes(workbook, sheet_name)

    summary = {
        "synced": 0, "skipped_duplicate": 0, "failed": 0,
        "not_found": 0, "total": len(call_notes),
        "not_found_contacts": [],  # list of (name, sheet_date) tuples
        "log_entries": [],  # list of SyncLogEntry
    }

    for i, cn in enumerate(call_notes, 1):
        logger.info(f"[{i}/{len(call_notes)}] Processing: {cn.name}")
        note_body = format_note_body(cn)

        if dry_run:
            logger.info(f"  DRY RUN - Would create note: '{note_body[:80]}...'")
            summary["synced"] += 1
            summary["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="synced (dry_run)", note_body=note_body,
            ))
            continue

        # Search for the contact in GHL
        contact = search_contact_by_name(cn.name)
        if not contact:
            summary["not_found"] += 1
            summary["not_found_contacts"].append((cn.name, cn.sheet_date))
            summary["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="not_found", note_body=note_body,
            ))
            continue

        contact_id = contact["id"]

        # Check for duplicates
        if note_already_exists(contact_id, note_body):
            logger.info(f"  Skipping duplicate note for '{cn.name}'")
            summary["skipped_duplicate"] += 1
            summary["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="skipped_duplicate", note_body=note_body,
                ghl_contact_id=contact_id,
            ))
            continue

        # Create the note
        result = create_note(contact_id, note_body)
        if result:
            summary["synced"] += 1
            summary["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="synced", note_body=note_body,
                ghl_contact_id=contact_id,
            ))
        else:
            summary["failed"] += 1
            summary["log_entries"].append(SyncLogEntry(
                name=cn.name, sheet_date=cn.sheet_date,
                status="failed", note_body=note_body,
                ghl_contact_id=contact_id,
            ))

        # Rate limit
        time.sleep(REQUEST_DELAY)

    _print_summary(summary, sheet_name)

    # Write per-sheet log when running standalone (not from sync_all_notes)
    if write_log and summary["log_entries"]:
        sheet_label = parse_sheet_date(sheet_name).replace(".", "-")
        write_sync_log(summary["log_entries"], label=sheet_label)

    return summary


def write_sync_log(log_entries: list, label: str = "sync") -> str:
    """Write sync log entries to a timestamped CSV file in sync_logs/."""
    os.makedirs(os.path.abspath(LOG_DIR), exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"sync_log_{label}_{timestamp}.csv"
    filepath = os.path.join(os.path.abspath(LOG_DIR), filename)

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Name", "Sheet Date", "Status", "GHL Contact ID", "Note Body", "Timestamp"])
        for entry in log_entries:
            writer.writerow([
                entry.name, entry.sheet_date, entry.status,
                entry.ghl_contact_id, entry.note_body, entry.timestamp,
            ])

    logger.info(f"Sync log written to: {filepath}")
    return filepath


def _print_summary(summary: dict, label: str) -> None:
    """Print a sync summary to the logger."""
    logger.info("=" * 60)
    logger.info(f"Sync Summary ({label}):")
    logger.info(f"  Total call notes:    {summary['total']}")
    logger.info(f"  Synced:              {summary['synced']}")
    logger.info(f"  Skipped (duplicate): {summary['skipped_duplicate']}")
    logger.info(f"  Not found in GHL:    {summary['not_found']}")
    logger.info(f"  Failed:              {summary['failed']}")
    if summary["not_found_contacts"]:
        logger.info("")
        logger.info("  Contacts NOT found in GHL:")
        for name, sheet_date in summary["not_found_contacts"]:
            logger.info(f"    - {name} (from sheet {sheet_date})")
    logger.info("=" * 60)


def sync_all_notes(dry_run: bool = False, excel_path: str | None = None) -> dict:
    """
    Sync call notes from ALL Economics sheets to GHL.

    Returns:
        Aggregated summary dict across all sheets.
    """
    _require_ghl_env()

    workbook = resolve_workbook_path(excel_path or EXCEL_PATH)
    sheets = get_economics_sheets_chronological(workbook)
    logger.info(f"Found {len(sheets)} Economics sheets to sync (oldest → newest)")

    totals = {
        "synced": 0, "skipped_duplicate": 0, "failed": 0,
        "not_found": 0, "total": 0,
        "not_found_contacts": [],
        "log_entries": [],
        "sheets_processed": 0,
    }

    for sheet_name in sheets:
        logger.info(f"\n{'#' * 60}")
        logger.info(f"Processing sheet: {sheet_name}")
        logger.info(f"{'#' * 60}")
        try:
            summary = sync_notes(
                sheet_name=sheet_name,
                dry_run=dry_run,
                write_log=False,
                excel_path=excel_path,
            )
            totals["synced"] += summary["synced"]
            totals["skipped_duplicate"] += summary["skipped_duplicate"]
            totals["failed"] += summary["failed"]
            totals["not_found"] += summary["not_found"]
            totals["total"] += summary["total"]
            totals["not_found_contacts"].extend(summary["not_found_contacts"])
            totals["log_entries"].extend(summary.get("log_entries", []))
            totals["sheets_processed"] += 1
        except Exception as e:
            logger.error(f"Error processing sheet '{sheet_name}': {e}")

    # Print consolidated summary with unique unmatched names
    logger.info(f"\n{'*' * 60}")
    logger.info("GRAND TOTAL - All Sheets:")
    logger.info(f"  Sheets processed:    {totals['sheets_processed']}/{len(sheets)}")
    logger.info(f"  Total call notes:    {totals['total']}")
    logger.info(f"  Synced:              {totals['synced']}")
    logger.info(f"  Skipped (duplicate): {totals['skipped_duplicate']}")
    logger.info(f"  Not found in GHL:    {totals['not_found']}")
    logger.info(f"  Failed:              {totals['failed']}")

    if totals["not_found_contacts"]:
        # Deduplicate names and show all sheets they appeared in
        not_found_map: dict[str, list[str]] = {}
        for name, sheet_date in totals["not_found_contacts"]:
            not_found_map.setdefault(name, []).append(sheet_date)

        logger.info("")
        logger.info(f"  Contacts NOT found in GHL ({len(not_found_map)} unique):")
        for name in sorted(not_found_map.keys()):
            dates = ", ".join(not_found_map[name])
            logger.info(f"    - {name}  (sheets: {dates})")

    logger.info(f"{'*' * 60}")

    # Write consolidated log CSV
    if totals["log_entries"]:
        write_sync_log(totals["log_entries"], label="all")

    return totals


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Sync call notes from Excel to GHL contact notes")
    parser.add_argument(
        "--file",
        type=str,
        default=None,
        help="Path to call notes workbook (default: GHL/jared_call_notes.xlsx)",
    )
    parser.add_argument(
        "--sheet",
        type=str,
        default=None,
        help="Name of a single Economics sheet to sync (e.g., '1.5.26 Economics')",
    )
    parser.add_argument(
        "--all",
        action="store_true",
        dest="sync_all",
        help="Sync all Economics sheets (oldest → newest)",
    )
    parser.add_argument(
        "--reset",
        action="store_true",
        help="Purge [Call Notes - ...] notes then re-sync all sheets oldest → newest",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Preview purge + sync without making API writes",
    )
    args = parser.parse_args()

    source = args.file or EXCEL_PATH
    if args.file:
        os.environ["CALL_NOTES_FILE"] = args.file

    if args.reset:
        reset_call_notes(excel_path=source, dry_run=args.dry_run)
    elif args.sync_all:
        sync_all_notes(dry_run=args.dry_run, excel_path=source)
    elif args.sheet:
        sync_notes(sheet_name=args.sheet, dry_run=args.dry_run, excel_path=source)
    else:
        parser.error("Please specify --reset, --sheet <name>, or --all")
